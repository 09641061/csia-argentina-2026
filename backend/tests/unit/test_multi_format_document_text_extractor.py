from __future__ import annotations

from io import BytesIO

import pymupdf
import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from PIL import Image

from app.analysis.domain.exceptions import AnalysisModelUnavailableError
from app.analysis.domain.model.valueobjects.vision_extraction_result import (
    VisionExtractionResult,
)
from app.analysis.infrastructure.text_extraction.multi_format_document_text_extractor import (
    MultiFormatDocumentTextExtractor,
)


class RecordingVisionClient:
    model_name = "fake-vision"

    def __init__(self) -> None:
        self.calls: list[tuple[str, str]] = []

    async def extract(
        self, *, image_content: bytes, mime_type: str, reference_label: str
    ) -> VisionExtractionResult:
        assert image_content
        self.calls.append((mime_type, reference_label))
        return VisionExtractionResult(
            visible_text="Correo visible: demo@example.com",
            visual_summary="Captura de una ficha de contacto.",
            document_type="contact_card",
            model_name=self.model_name,
        )


class UnavailableVisionClient(RecordingVisionClient):
    async def extract(self, **kwargs) -> VisionExtractionResult:
        del kwargs
        raise AnalysisModelUnavailableError("vision model unavailable")


def _pdf_bytes(text: str | None = None) -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def _docx_bytes() -> bytes:
    document = WordDocument()
    document.add_paragraph("Informe académico de la hackatón")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Responsable"
    table.cell(0, 1).text = "demo@example.com"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventario"
    sheet["A1"] = "Servicio"
    sheet["B1"] = "Costo"
    sheet["A2"] = "Sentinel"
    sheet["B2"] = "=40+2"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _png_bytes() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (32, 32), color="white").save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_extracts_text_pdf_without_calling_vision() -> None:
    vision = RecordingVisionClient()
    extractor = MultiFormatDocumentTextExtractor(vision)

    result = await extractor.extract_content(
        _pdf_bytes("Public project summary"), "application/pdf", "summary.pdf"
    )

    assert result["format"] == "pdf"
    assert result["pages"][0]["source"] == "embedded_text"
    assert "Public project summary" in result["pages"][0]["text"]
    assert vision.calls == []


@pytest.mark.asyncio
async def test_scanned_pdf_page_requires_local_vision() -> None:
    vision = RecordingVisionClient()
    extractor = MultiFormatDocumentTextExtractor(vision)

    result = await extractor.extract_content(
        _pdf_bytes(), "application/pdf", "scanned.pdf"
    )

    assert result["pages"][0]["source"] == "ollama_vision"
    assert result["pages"][0]["content"]["vision_model"] == "fake-vision"
    assert vision.calls == [("image/png", "scanned.pdf, page 1")]


@pytest.mark.asyncio
async def test_extracts_word_paragraphs_and_tables() -> None:
    extractor = MultiFormatDocumentTextExtractor(RecordingVisionClient())

    result = await extractor.extract_content(
        _docx_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "report.docx",
    )

    assert result["format"] == "docx"
    assert result["paragraphs"][0]["text"] == "Informe académico de la hackatón"
    assert result["tables"][0]["cells"][1]["text"] == "demo@example.com"


@pytest.mark.asyncio
async def test_extracts_excel_values_and_preserves_formulas_as_text() -> None:
    extractor = MultiFormatDocumentTextExtractor(RecordingVisionClient())

    result = await extractor.extract_content(
        _xlsx_bytes(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "inventory.xlsx",
    )

    assert result["format"] == "xlsx"
    cells = result["sheets"][0]["cells"]
    assert {cell["coordinate"]: cell["value"] for cell in cells}["B2"] == "=40+2"


@pytest.mark.asyncio
async def test_image_content_comes_from_mandatory_local_vision() -> None:
    vision = RecordingVisionClient()
    extractor = MultiFormatDocumentTextExtractor(vision)

    result = await extractor.extract_content(_png_bytes(), "image/png", "contact.png")

    assert result["visible_text"] == "Correo visible: demo@example.com"
    assert result["vision_model"] == "fake-vision"
    assert vision.calls == [("image/png", "contact.png")]


@pytest.mark.asyncio
async def test_image_is_not_analyzable_when_local_vision_is_unavailable() -> None:
    extractor = MultiFormatDocumentTextExtractor(UnavailableVisionClient())

    with pytest.raises(AnalysisModelUnavailableError):
        await extractor.extract_content(_png_bytes(), "image/png", "contact.png")
