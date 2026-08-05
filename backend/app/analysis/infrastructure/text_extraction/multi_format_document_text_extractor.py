from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

import pymupdf
from docx import Document as WordDocument
from openpyxl import load_workbook

from app.analysis.application.internal.outboundservices.document_text_extractor import (
    DocumentTextExtractor,
)
from app.analysis.application.internal.outboundservices.ollama_vision_extraction_client import (
    OllamaVisionExtractionClient,
)
from app.analysis.domain.exceptions import DocumentContentExtractionError
from app.analysis.domain.model.valueobjects.json_types import JsonContainer, JsonValue
from app.analysis.infrastructure.text_extraction.json_document_text_extractor import (
    JsonDocumentTextExtractor,
)
from app.documents.domain.model.valueobjects.document_mime_type import DocumentMimeType

_MAX_EXTRACTED_CHARACTERS = 200_000
_MAX_PDF_PAGES = 50
_MAX_SCANNED_PDF_PAGES = 10
_MAX_DOCX_PARAGRAPHS = 5_000
_MAX_DOCX_TABLE_CELLS = 20_000
_MAX_XLSX_SHEETS = 50
_MAX_XLSX_NON_EMPTY_CELLS = 50_000
_MAX_XLSX_SCANNED_CELLS = 250_000
_MAX_CELL_CHARACTERS = 5_000
_MAX_RENDERED_PAGE_PIXELS = 20_000_000


class MultiFormatDocumentTextExtractor(DocumentTextExtractor):
    """Turns supported files into a bounded JSON-like representation for the AI."""

    def __init__(self, vision_client: OllamaVisionExtractionClient) -> None:
        self._json_extractor = JsonDocumentTextExtractor()
        self._vision_client = vision_client

    async def extract_content(
        self, content: bytes, mime_type: str, original_filename: str
    ) -> JsonContainer:
        normalized_mime_type = DocumentMimeType(mime_type).value

        if normalized_mime_type == DocumentMimeType.JSON:
            extracted = await self._json_extractor.extract_content(
                content, normalized_mime_type, original_filename
            )
        elif normalized_mime_type == DocumentMimeType.PDF:
            extracted = await self._extract_pdf(content, original_filename)
        elif normalized_mime_type == DocumentMimeType.DOCX:
            extracted = self._extract_docx(content, original_filename)
        elif normalized_mime_type == DocumentMimeType.XLSX:
            extracted = self._extract_xlsx(content, original_filename)
        elif normalized_mime_type in {DocumentMimeType.PNG, DocumentMimeType.JPEG}:
            extracted = await self._extract_image(
                content, normalized_mime_type, original_filename
            )
        else:  # pragma: no cover - DocumentMimeType rejects this first
            raise DocumentContentExtractionError(
                f"Unsupported document type for {original_filename}"
            )

        if self._character_count(extracted) > _MAX_EXTRACTED_CHARACTERS:
            raise DocumentContentExtractionError(
                f"Unable to analyze {original_filename}: extracted content exceeds "
                f"the {_MAX_EXTRACTED_CHARACTERS}-character safety limit"
            )
        return extracted

    async def _extract_pdf(
        self, content: bytes, original_filename: str
    ) -> JsonContainer:
        try:
            document = pymupdf.open(stream=content, filetype="pdf")
        except Exception as error:
            raise DocumentContentExtractionError(
                f"Unable to read PDF content from {original_filename}"
            ) from error

        try:
            if document.needs_pass:
                raise DocumentContentExtractionError(
                    f"Unable to analyze {original_filename}: password-protected PDFs "
                    "are not supported"
                )
            if document.page_count == 0:
                raise DocumentContentExtractionError(
                    f"Unable to analyze {original_filename}: the PDF has no pages"
                )
            if document.page_count > _MAX_PDF_PAGES:
                raise DocumentContentExtractionError(
                    f"Unable to analyze {original_filename}: PDFs are limited to "
                    f"{_MAX_PDF_PAGES} pages"
                )

            pages: list[JsonValue] = []
            scanned_page_count = 0
            for page_index, page in enumerate(document, start=1):
                try:
                    embedded_text = self._clean_text(page.get_text("text"))
                except Exception as error:
                    raise DocumentContentExtractionError(
                        f"Unable to extract page {page_index} from {original_filename}"
                    ) from error

                if embedded_text:
                    pages.append(
                        {
                            "page": page_index,
                            "source": "embedded_text",
                            "text": embedded_text,
                        }
                    )
                    continue

                scanned_page_count += 1
                if scanned_page_count > _MAX_SCANNED_PDF_PAGES:
                    raise DocumentContentExtractionError(
                        f"Unable to analyze {original_filename}: scanned PDFs are limited "
                        f"to {_MAX_SCANNED_PDF_PAGES} image-only pages"
                    )
                rendered_width = max(1, int(page.rect.width * 1.5))
                rendered_height = max(1, int(page.rect.height * 1.5))
                if rendered_width * rendered_height > _MAX_RENDERED_PAGE_PIXELS:
                    raise DocumentContentExtractionError(
                        f"Unable to render page {page_index} from {original_filename}: "
                        "page dimensions exceed the safe visual-processing limit"
                    )
                try:
                    pixmap = page.get_pixmap(
                        matrix=pymupdf.Matrix(1.5, 1.5), alpha=False
                    )
                    rendered_page = pixmap.tobytes("png")
                except Exception as error:
                    raise DocumentContentExtractionError(
                        f"Unable to render page {page_index} from {original_filename}"
                    ) from error

                vision_result = await self._vision_client.extract(
                    image_content=rendered_page,
                    mime_type=DocumentMimeType.PNG,
                    reference_label=f"{original_filename}, page {page_index}",
                )
                pages.append(
                    {
                        "page": page_index,
                        "source": "ollama_vision",
                        "content": vision_result.to_document_payload(),
                    }
                )

            return {
                "format": "pdf",
                "page_count": document.page_count,
                "image_only_pages": scanned_page_count,
                "pages": pages,
            }
        finally:
            document.close()

    def _extract_docx(
        self, content: bytes, original_filename: str
    ) -> JsonContainer:
        try:
            document = WordDocument(BytesIO(content))
        except Exception as error:
            raise DocumentContentExtractionError(
                f"Unable to read Word content from {original_filename}"
            ) from error

        paragraphs: list[JsonValue] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            if index > _MAX_DOCX_PARAGRAPHS:
                raise DocumentContentExtractionError(
                    f"Unable to analyze {original_filename}: Word documents are limited "
                    f"to {_MAX_DOCX_PARAGRAPHS} paragraphs"
                )
            text = self._clean_text(paragraph.text)
            if text:
                paragraphs.append({"paragraph": index, "text": text})

        tables: list[JsonValue] = []
        table_cell_count = 0
        for table_index, table in enumerate(document.tables, start=1):
            cells: list[JsonValue] = []
            for row_index, row in enumerate(table.rows, start=1):
                for column_index, cell in enumerate(row.cells, start=1):
                    table_cell_count += 1
                    if table_cell_count > _MAX_DOCX_TABLE_CELLS:
                        raise DocumentContentExtractionError(
                            f"Unable to analyze {original_filename}: Word tables are "
                            f"limited to {_MAX_DOCX_TABLE_CELLS} cells"
                        )
                    text = self._clean_text(cell.text)
                    if text:
                        cells.append(
                            {
                                "row": row_index,
                                "column": column_index,
                                "text": text,
                            }
                        )
            if cells:
                tables.append({"table": table_index, "cells": cells})

        headers_and_footers: list[JsonValue] = []
        seen_section_text: set[tuple[str, str]] = set()
        for section_index, section in enumerate(document.sections, start=1):
            for area_name, area in (
                ("header", section.header),
                ("footer", section.footer),
            ):
                text = self._clean_text(
                    "\n".join(paragraph.text for paragraph in area.paragraphs)
                )
                deduplication_key = (area_name, text)
                if text and deduplication_key not in seen_section_text:
                    seen_section_text.add(deduplication_key)
                    headers_and_footers.append(
                        {"section": section_index, "area": area_name, "text": text}
                    )

        if not paragraphs and not tables and not headers_and_footers:
            raise DocumentContentExtractionError(
                f"Unable to analyze {original_filename}: no readable Word content was found"
            )
        return {
            "format": "docx",
            "paragraphs": paragraphs,
            "tables": tables,
            "headers_and_footers": headers_and_footers,
        }

    def _extract_xlsx(
        self, content: bytes, original_filename: str
    ) -> JsonContainer:
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
        except Exception as error:
            raise DocumentContentExtractionError(
                f"Unable to read Excel content from {original_filename}"
            ) from error

        try:
            if len(workbook.worksheets) > _MAX_XLSX_SHEETS:
                raise DocumentContentExtractionError(
                    f"Unable to analyze {original_filename}: workbooks are limited to "
                    f"{_MAX_XLSX_SHEETS} sheets"
                )

            sheets: list[JsonValue] = []
            non_empty_cell_count = 0
            scanned_cell_count = 0
            for worksheet in workbook.worksheets:
                cells: list[JsonValue] = []
                for row in worksheet.iter_rows():
                    for cell in row:
                        scanned_cell_count += 1
                        if scanned_cell_count > _MAX_XLSX_SCANNED_CELLS:
                            raise DocumentContentExtractionError(
                                f"Unable to analyze {original_filename}: the workbook's "
                                "declared cell range is too large"
                            )
                        if cell.value is None:
                            continue
                        non_empty_cell_count += 1
                        if non_empty_cell_count > _MAX_XLSX_NON_EMPTY_CELLS:
                            raise DocumentContentExtractionError(
                                f"Unable to analyze {original_filename}: workbooks are "
                                f"limited to {_MAX_XLSX_NON_EMPTY_CELLS} non-empty cells"
                            )
                        cells.append(
                            {
                                "coordinate": cell.coordinate,
                                "value": self._spreadsheet_value(cell.value),
                                "data_type": str(cell.data_type),
                            }
                        )
                sheets.append(
                    {
                        "name": self._bounded_text(worksheet.title),
                        "state": str(worksheet.sheet_state),
                        "cells": cells,
                    }
                )

            if non_empty_cell_count == 0:
                raise DocumentContentExtractionError(
                    f"Unable to analyze {original_filename}: no populated Excel cells were found"
                )
            return {
                "format": "xlsx",
                "sheet_count": len(sheets),
                "sheets": sheets,
            }
        finally:
            workbook.close()

    async def _extract_image(
        self, content: bytes, mime_type: str, original_filename: str
    ) -> JsonContainer:
        result = await self._vision_client.extract(
            image_content=content,
            mime_type=mime_type,
            reference_label=original_filename,
        )
        return result.to_document_payload()

    @classmethod
    def _spreadsheet_value(cls, value: Any) -> JsonValue:
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        if value is None or isinstance(value, (bool, int, float)):
            return value
        return cls._bounded_text(str(value))

    @staticmethod
    def _clean_text(value: str) -> str:
        return value.replace("\x00", "").strip()

    @staticmethod
    def _bounded_text(value: str) -> str:
        cleaned = value.replace("\x00", "").strip()
        if len(cleaned) > _MAX_CELL_CHARACTERS:
            return f"{cleaned[:_MAX_CELL_CHARACTERS]}...[truncated]"
        return cleaned

    @classmethod
    def _character_count(cls, value: JsonValue) -> int:
        if isinstance(value, str):
            return len(value)
        if isinstance(value, list):
            return sum(cls._character_count(item) for item in value)
        if isinstance(value, dict):
            return sum(
                len(key) + cls._character_count(item) for key, item in value.items()
            )
        return 0
